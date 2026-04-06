from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form
from typing import Optional, List
from pydantic import BaseModel
import os

from core.auth import get_current_user
from db.models import User
from integrations.google_drive import (
    list_drive_folders, 
    create_drive_folder, 
    upload_to_drive,
    get_drive_service
)

router = APIRouter()

class CreateFolderRequest(BaseModel):
    folder_name: str
    parent_id: Optional[str] = None

@router.get("/folders")
async def get_drive_folders(user: User = Depends(get_current_user)):
    """Mendapatkan daftar folder Google Drive untuk tree navigation."""
    service = get_drive_service()
    if not service:
        raise HTTPException(400, "Google Drive Credentials belum diatur di Integrations.")
        
    try:
        folders = await list_drive_folders()
        return {"status": "success", "folders": folders}
    except Exception as e:
        raise HTTPException(500, str(e))

@router.post("/folders")
async def api_create_drive_folder(req: CreateFolderRequest, user: User = Depends(get_current_user)):
    """Membuat folder baru di Google Drive."""
    service = get_drive_service()
    if not service:
        raise HTTPException(400, "Google Drive Credentials belum diatur.")
        
    try:
        folder_id = await create_drive_folder(req.folder_name, req.parent_id)
        return {"status": "success", "folder_id": folder_id}
    except Exception as e:
        raise HTTPException(500, str(e))

@router.post("/upload")
async def api_upload_to_drive(
    file: UploadFile = File(...),
    folder_id: Optional[str] = Form(None),
    user: User = Depends(get_current_user)
):
    """Mengupload file yang dibuat AI ke Google Drive."""
    service = get_drive_service()
    if not service:
        raise HTTPException(400, "Google Drive Credentials belum diatur.")
        
    try:
        content = await file.read()
        res = await upload_to_drive(file.filename, content, folder_id)
        
        if res.get("status") == "error":
            raise HTTPException(500, res.get("message"))
            
        return res
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(500, f"Upload error: {str(e)}")

@router.post("/upload_generated")
async def api_upload_generated_to_drive(
    content: str = Form(...),
    format: str = Form(...),
    filename: str = Form(...),
    folder_id: Optional[str] = Form(None),
    user: User = Depends(get_current_user)
):
    """Konversi text dari AI menjadi file fisik lalu upload ke Drive."""
    service = get_drive_service()
    if not service:
        raise HTTPException(400, "Google Drive Credentials belum diatur.")
        
    import io
    file_bytes = b""
    format = format.lower()
    
    try:
        if format in ["txt", "csv", "md"]:
            file_bytes = content.encode("utf-8")
            
        elif format == "docx":
            from docx import Document
            doc = Document()
            for line in content.split("\n"):
                doc.add_paragraph(line)
            buf = io.BytesIO()
            doc.save(buf)
            file_bytes = buf.getvalue()
            
        elif format == "xlsx":
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            for row in content.strip().split("\n"):
                # Simple CSV parsing for XLSX
                ws.append(row.split(",")) 
            buf = io.BytesIO()
            wb.save(buf)
            file_bytes = buf.getvalue()
            
        elif format == "pdf":
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Paragraph
            from reportlab.lib.styles import getSampleStyleSheet
            buf = io.BytesIO()
            doc = SimpleDocTemplate(buf, pagesize=letter)
            styles = getSampleStyleSheet()
            flowables = []
            for line in content.split("\n"):
                if line.strip():
                    # Sanitize line for reportlab XML
                    sanitized = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
                    flowables.append(Paragraph(sanitized, styles["Normal"]))
            doc.build(flowables)
            file_bytes = buf.getvalue()
        else:
            raise HTTPException(400, f"Format tidak didukung: {format}")
            
        res = await upload_to_drive(filename, file_bytes, folder_id)
        if res.get("status") == "error":
            raise HTTPException(500, res.get("message"))
        return res
        
    except HTTPException:
        raise
    except Exception as e:
        import traceback
        traceback.print_exc()
        raise HTTPException(500, f"Konversi data ke file fisik gagal: {str(e)}")
